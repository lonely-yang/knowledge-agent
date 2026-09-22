import io
import asyncio
from datetime import datetime
from docx import Document
from crud.RustFS import settings

# 命名空间常量
NS_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

async def extract_docx_text_and_images(docx_bytes, name_part, s3):
    bio = io.BytesIO(docx_bytes)
    doc = Document(bio)
    # 文档标题
    docx_title = doc.core_properties.title or name_part
    result = []
    all_images = []
    img_idx_global = 1

    try:
        # 遍历文档body内所有块（段落、表格，顺序保持word原文）
        for block in doc.element.body:
            # 段落
            if block.tag.endswith("p"):
                para = next((p for p in doc.paragraphs if p._element == block), None)
                if not para:
                    continue
                para_text = para.text.strip()
                para_image_links = []

                # 遍历run提取内联图片
                for run in para.runs:
                    # 使用findall + 完整命名空间，替代xpath+namespaces
                    blips = run._r.findall(f".//{NS_A}blip")
                    for blip in blips:
                        r_embed = blip.get(f"{NS_R}embed")
                        if not r_embed:
                            continue
                        rel = doc.part.rels.get(r_embed)
                        if rel is None:
                            continue
                        img_part = rel.target_part
                        img_bytes = img_part.blob
                        ext = img_part.content_type.split("/")[-1]

                        # 图片key命名
                        ts = int(datetime.utcnow().timestamp())
                        img_key = f"{name_part}-para{len(result)+1}-img{img_idx_global}-{ts}.{ext}"

                        # 上传RustFS
                        await asyncio.to_thread(
                            s3.put_object,
                            Bucket=settings.RUSTFS_BUCKET,
                            Key=img_key,
                            Body=img_bytes,
                            ContentType=img_part.content_type
                        )
                        # 预签名url
                        preview_url = await asyncio.to_thread(
                            s3.generate_presigned_url,
                            ClientMethod="get_object",
                            Params={"Bucket": settings.RUSTFS_BUCKET, "Key": img_key},
                            ExpiresIn=settings.PRESIGNED_EXPIRE
                        )
                        img_info = {
                            "key": img_key,
                            "preview_url": preview_url,
                            "ext": ext
                        }
                        all_images.append(img_info)
                        para_image_links.append(f"\n[IMAGE{img_idx_global}] {preview_url}")
                        img_idx_global += 1

                # 图片链接追加到段落文本末尾
                if para_image_links:
                    para_text += "\n=====图片预览链接====="
                    para_text += "".join(para_image_links)

                result.append({
                    "type": "paragraph",
                    "text": para_text
                })

            # 表格
            elif block.tag.endswith("tbl"):
                table = next((t for t in doc.tables if t._element == block), None)
                if not table:
                    continue
                table_lines = []
                for row in table.rows:
                    cell_texts = []
                    for cell in row.cells:
                        cell_texts.append(cell.text.strip())
                    table_lines.append(" | ".join(cell_texts))
                table_text = "\n".join(table_lines)
                result.append({
                    "type": "table",
                    "text": table_text
                })
    finally:
        bio.close()

    text = ''
    for res in result:
        text = text + res.get('text')
    return {
        "title": docx_title,
        "text": text,
        "images": all_images
    }
