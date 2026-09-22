import io
import asyncio
from datetime import datetime
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from core.config import settings

async def extract_pptx_text_and_images(pptx_bytes, name_part, s3):
    bio = io.BytesIO(pptx_bytes)
    prs = Presentation(bio)
    # 文档标题，优先core_properties，没有就用文件名
    pptx_title = prs.core_properties.title or name_part
    result = []
    all_images = []
    img_idx_global = 1

    # 递归遍历形状，处理组合形状，提取图片和文本
    async def iterate_shape(shape, slide_num, slide_text_parts, slide_image_links):
        nonlocal img_idx_global
        # 组合形状：递归展开
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for sub_shape in shape.shapes:
                await iterate_shape(sub_shape, slide_num, slide_text_parts, slide_image_links)
            return
        # 图片
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img = shape.image
            img_bytes = img.blob
            ext = img.content_type.split("/")[-1]
            ts = int(datetime.utcnow().timestamp())
            img_key = f"{name_part}-slide{slide_num}-img{img_idx_global}-{ts}.{ext}"

            # 上传RustFS
            await asyncio.to_thread(
                s3.put_object,
                Bucket=settings.RUSTFS_BUCKET,
                Key=img_key,
                Body=img_bytes,
                ContentType=img.content_type
            )
            preview_url = await asyncio.to_thread(
                s3.generate_presigned_url,
                ClientMethod="get_object",
                Params={"Bucket": settings.RUSTFS_BUCKET, "Key": img_key},
                ExpiresIn=settings.PRESIGNED_EXPIRE
            )
            img_info = {
                "key": img_key,
                "preview_url": preview_url,
                "ext": ext,
                "slide": slide_num
            }
            all_images.append(img_info)
            slide_image_links.append(f"\n[IMAGE{img_idx_global}] {preview_url}")
            img_idx_global +=1
            return
        # 文本框
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                text = para.text.strip()
                if text:
                    slide_text_parts.append(text)
        # 表格
        if shape.has_table:
            table_lines = []
            for row in shape.table.rows:
                cell_texts = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        cell_texts.append(cell_text)
                table_lines.append(" | ".join(cell_texts))
            table_block = "\n".join(table_lines)
            slide_text_parts.append(table_block)

    try:
        # 遍历所有幻灯片
        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx +1
            slide_text_parts = []
            slide_image_links = []

            # 遍历页面所有形状（含组合图），增加 await
            for shape in slide.shapes:
                await iterate_shape(shape, slide_num, slide_text_parts, slide_image_links)

            slide_text = "\n".join(slide_text_parts)
            # 追加图片链接到本页文本末尾
            if slide_image_links:
                slide_text += "\n=====图片预览链接====="
                slide_text += "".join(slide_image_links)

            result.append({
                "slide": slide_num,
                "text": slide_text
            })
    finally:
        bio.close()
    text = ''
    for res in result:
        text = text + res.get('text')
    return {
        "title": pptx_title,
        "text": text,
        "images": all_images
    }
